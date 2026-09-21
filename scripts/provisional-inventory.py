"""Provisional inventory candidate generator.

Reads the Owner-supplied provisional workbook and produces the audit figures, a per-row
catalogue mapping and a V3-shaped candidate file. Nothing here is import input: the source is
provisional, so the candidate is deliberately marked and cannot be staged.

Every mapping is deterministic. A value that is not present in the source is never invented;
it is reported as an Owner decision or as a blocking reason.
"""
import csv, hashlib, re, sys, collections
try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl required')

EXPECTED_SHA256 = '5ff5ce67c8ec1afbfab6bc5beeab0a2a725b73db0e1958a02860a0db39fbe1f4'
V3_HEADER = ['source_kind','intent','model_id','season','variant_id','manufacturer_sku','quantity','unit','asset_ids','store_id','source_document','source_row','category','size','tier','bsl_mm','status','manufacturer','model_name','note']

FAMILY = {
    '大人スキー':('SKI','ADULT'), '子供スキー':('SKI','KIDS'),
    '大人スキーブーツ':('SKI_BOOT','ADULT'), '子供スキーブーツ':('SKI_BOOT','KIDS'),
    'スキーポール':('POLE','ADULT'),
    '大人ボード':('SNOWBOARD','ADULT'), '子供ボード':('SNOWBOARD','KIDS'), 'キッズボード':('SNOWBOARD','KIDS'),
    '大人ボードブーツ':('SNOWBOARD_BOOT','ADULT'), '子供ボードブーツ':('SNOWBOARD_BOOT','KIDS'), 'キッズボードブーツ':('SNOWBOARD_BOOT','KIDS'),
    'ヘルメット':('HELMET','ADULT'),
    'ボードバイン':('SNOWBOARD_BINDING','ADULT'),
}
# Owner decision for this round: only these four families are registered now. A snowboard is
# one lending unit — board plus its mounted binding — so the binding never becomes an Asset,
# QR label, price, reservation stock or quantity pool of its own. Boots are one Asset per
# left/right pair; both sides carry the same Asset ID when both are labelled.
IN_SCOPE = {'SKI':'ASSET_PAIR','SNOWBOARD':'ASSET_BOARD','SKI_BOOT':'ASSET_PAIR','SNOWBOARD_BOOT':'ASSET_PAIR'}
# Excluded by Owner scope for this round. Not data defects, not unknown families, not
# unclassified rows: the catalogue, importer, recommendation, reservation and label support for
# them is untouched and they keep their source rows and quantities.
EXCLUDED_BY_OWNER_SCOPE = {'SNOWBOARD_BINDING','HELMET','POLE'}
# Added once the Owner supplies inventory material for them; absent from this source.
FUTURE_INPUT_REQUIRED = {'WEAR_JACKET','WEAR_PANTS'}
UNIT = dict(IN_SCOPE)
ASSET_NOTE = {
 'SNOWBOARD':'BOARD_WITH_MOUNTED_BINDING_ONE_ASSET',
 'SKI':'ONE_PAIR_ONE_ASSET',
 'SKI_BOOT':'LEFT_RIGHT_PAIR_ONE_ASSET',
 'SNOWBOARD_BOOT':'LEFT_RIGHT_PAIR_ONE_ASSET',
}

def digest(path):
    h = hashlib.sha256()
    with open(path,'rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()

def read(path):
    ws = openpyxl.load_workbook(path, data_only=True)['Sheet1']
    rows = []
    for number, values in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if all(v is None or str(v).strip() == '' for v in values):
            continue
        rows.append({'row':number,'destination':values[0],'destination_name':values[1],'comment':values[2],
                     'code':str(values[3]).strip(),'name':str(values[4]).strip(),
                     'jp_size':('' if values[5] is None else str(values[5]).strip()),
                     'quantity':int(values[6]),'category':str(values[8]).strip(),'hierarchy':str(values[9]).strip()})
    return rows, ws['G1'].value

def season(comment):
    # The only season marker present is the 2627 token in the shipment comment.
    return '2026/27' if re.search(r'2627', str(comment or '')) else ''

def ski_size(row):
    """Skis carry no JP size. The length is the trailing number of the product name and is
    cross-checked against the trailing digits of the item code; disagreement blocks the row."""
    name_match = re.search(r'(\d{2,3})\s*$', row['name'])
    code_match = re.search(r'(\d{3})$', row['code'])
    if not name_match or not code_match:
        return '', 'SKI_LENGTH_NOT_DETERMINABLE'
    if name_match.group(1).lstrip('0') != code_match.group(1).lstrip('0'):
        return '', 'SKI_LENGTH_NAME_CODE_MISMATCH'
    return name_match.group(1) + ' cm', ''

def classify(row):
    mapped = FAMILY.get(row['category'])
    blocking = []
    if not mapped:
        return {'family':'UNMAPPED_CATEGORY','age':'','size':'','blocking':['UNKNOWN_SOURCE_CATEGORY']}
    family, age = mapped
    brand = 'SALOMON' if row['hierarchy'].replace(' ','').startswith('SAL') else ''
    if not brand:
        blocking.append('BRAND_NOT_DETERMINABLE')
    if not season(row['comment']):
        blocking.append('SEASON_NOT_DETERMINABLE')
    if family in EXCLUDED_BY_OWNER_SCOPE:
        blocking.append('EXCLUDED_BY_OWNER_SCOPE')
        size = row['jp_size']
    elif family == 'SKI':
        size, issue = ski_size(row)
        if issue:
            blocking.append(issue)
    else:
        size = row['jp_size']
        if not size:
            blocking.append('SIZE_NOT_PRESENT')
    # Present in no form anywhere in the source.
    blocking.append('TIER_OWNER_MAPPING_REQUIRED')
    blocking.append('STORE_OWNER_ALLOCATION_REQUIRED')
    scope = 'IN_SCOPE' if family in IN_SCOPE else ('EXCLUDED_BY_OWNER_SCOPE' if family in EXCLUDED_BY_OWNER_SCOPE else 'UNCLASSIFIED')
    return {'family':family,'age':age,'size':size,'brand':brand,'blocking':blocking,'scope':scope}

def main(path, out_dir):
    actual = digest(path)
    if actual != EXPECTED_SHA256:
        sys.exit(f'SOURCE_HASH_MISMATCH expected={EXPECTED_SHA256} actual={actual}')
    rows, declared_total = read(path)
    mapping_path = out_dir + '/PROVISIONAL_CATALOG_MAPPING.csv'
    candidate_path = out_dir + '/V3_CANDIDATE_NOT_FOR_IMPORT.csv'
    with open(mapping_path,'w',newline='',encoding='utf8') as fh:
        w = csv.writer(fh, lineterminator='\n')
        w.writerow(['source_row','item_code','source_name','source_category','source_jp_size','quantity',
                    'mapped_family','owner_scope','asset_model','age','candidate_season','candidate_brand','candidate_size',
                    'tier_status','store_status','bsl_status','importable','blocking_reason'])
        for row in rows:
            c = classify(row)
            supported = c['family'] in UNIT
            bsl = 'UNVERIFIED_NULL' if c['family'] == 'SKI_BOOT' else 'NOT_APPLICABLE'
            w.writerow([row['row'],row['code'],row['name'],row['category'],row['jp_size'],row['quantity'],
                        c['family'],c['scope'],ASSET_NOTE.get(c['family'],'NOT_REGISTERED_THIS_ROUND'),
                        c['age'],season(row['comment']),c.get('brand',''),c['size'],
                        'OWNER_MAPPING_REQUIRED','OWNER_ALLOCATION_REQUIRED',bsl,
                        'NO' if not supported else 'PENDING_OWNER_DECISIONS','|'.join(c['blocking'])])
    with open(candidate_path,'w',newline='',encoding='utf8') as fh:
        w = csv.writer(fh, lineterminator='\n')
        w.writerow(V3_HEADER)
        for row in rows:
            c = classify(row)
            if c['family'] not in UNIT:
                continue
            unit = UNIT[c['family']]
            w.writerow(['SHOP_RECEIPT','ADD','OWNER_REQUIRED_MODEL_ID',season(row['comment']),
                        'OWNER_REQUIRED_VARIANT_ID','',row['quantity'],unit,
                        'OWNER_REQUIRED_ASSET_IDS' if unit.startswith('ASSET_') else '',
                        'OWNER_REQUIRED_STORE','【(株)Yuge 山形蔵王】新店舗投入予定明細_20260914.xlsx',
                        f"row{row['row']}",c['family'],c['size'],'OWNER_REQUIRED_TIER','', 'UNVERIFIED',
                        c.get('brand',''),row['name'],ASSET_NOTE[c['family']]])
    categories = collections.Counter()
    for row in rows:
        categories[row['category']] += row['quantity']
    families = collections.Counter()
    for row in rows:
        families[classify(row)['family']] += row['quantity']
    print('sourceSha256', actual)
    print('detailRows', len(rows))
    print('declaredTotal', declared_total, 'summedTotal', sum(r['quantity'] for r in rows))
    print('byCategory', dict(categories))
    print('byFamily', dict(families))
    scopes = collections.Counter()
    scope_rows = collections.Counter()
    for row in rows:
        c = classify(row)
        scopes[c['scope']] += row['quantity']
        scope_rows[c['scope']] += 1
    print('SOURCE_RAW_TOTAL', sum(r['quantity'] for r in rows), 'rows', len(rows))
    print('IN_SCOPE_TOTAL', scopes['IN_SCOPE'], 'rows', scope_rows['IN_SCOPE'])
    print('EXCLUDED_BY_OWNER_SCOPE', scopes['EXCLUDED_BY_OWNER_SCOPE'], 'rows', scope_rows['EXCLUDED_BY_OWNER_SCOPE'])
    print('FUTURE_INPUT_REQUIRED', sorted(FUTURE_INPUT_REQUIRED), 'presentInSource', 0)
    print('candidateRows', sum(1 for r in rows if classify(r)['family'] in IN_SCOPE))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
